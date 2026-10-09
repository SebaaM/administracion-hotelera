import re
from io import BytesIO
from zipfile import ZipFile
from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.styles.colors import COLOR_INDEX
from rest_framework.exceptions import ValidationError

LIMIT=5*1024*1024
MONTHS={"ene":1,"feb":2,"mar":3,"abr":4,"may":5,"jun":6,"jul":7,"ago":8,"sep":9,"oct":10,"nov":11,"dic":12}

def label(value):
    if isinstance(value,(float,int)) and not isinstance(value,bool) and float(value).is_integer():
        return str(int(value))
    return str(value or "").strip()

def checked_range(reference):
    if not isinstance(reference,str) or len(reference)>40:raise ValueError("referencia")
    left,top,right,bottom=range_boundaries(reference)
    if any(v is None for v in (left,top,right,bottom)) or not (1<=left<=right<=400 and 1<=top<=bottom<=5000):
        raise ValidationError("El Excel supera 5.000 filas o 400 columnas.")
    return (right-left+1)*(bottom-top+1)

def check_archive(content):
    if not content or len(content)>LIMIT:
        raise ValidationError("El archivo debe tener contenido y no superar 5 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            names=archive.namelist()
            if len(names)!=len(set(names)) or len(names)>1000:
                raise ValidationError("El Excel tiene una estructura inválida.")
            if any("vba" in name.lower() for name in names):
                raise ValidationError("No se admiten macros.")
            if sum(i.file_size for i in archive.infolist())>32*1024*1024:
                raise ValidationError("El Excel descomprimido supera 32 MB.")
            total=expanded=comments=worksheets=0
            for name in names:
                if not name.endswith(".xml"):continue
                root=fromstring(archive.read(name),forbid_dtd=True,forbid_entities=True)
                worksheet=root.tag.rsplit("}",1)[-1]=="worksheet"
                if worksheet:
                    worksheets+=1
                    if worksheets>20:raise ValidationError("Se admiten hasta 20 hojas.")
                for element in root.iter():
                    tag=element.tag.rsplit("}",1)[-1]
                    if worksheet:
                        if tag=="dimension":checked_range(element.get("ref"))
                        elif tag=="row":
                            if not 1<=int(element.get("r","0"))<=5000:raise ValueError("fila")
                        elif tag=="c":
                            total+=1
                            if checked_range(element.get("r"))!=1:raise ValueError("celda")
                        elif tag in ("mergeCell","hyperlink"):
                            expanded+=checked_range(element.get("ref"))
                    if tag in ("comment","threadedComment"):
                        comments+=1
                        if checked_range(element.get("ref"))!=1:raise ValueError("nota")
                    if total>100000 or expanded>100000 or comments>100000:
                        raise ValidationError("El Excel supera 100.000 celdas o posiciones expandidas de rangos.")
    except ValidationError:raise
    except Exception as exc:
        raise ValidationError("Subí un Excel .xlsx válido, sin macros ni entidades XML.") from exc

def theme_colors(book):
    if not book.loaded_theme:return []
    root=fromstring(book.loaded_theme)
    ns={"a":"http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme=root.find(".//a:clrScheme",ns)
    if scheme is None:return []
    by_name={e.tag.rsplit("}",1)[-1]:list(e)[0].get("lastClr",list(e)[0].get("val","")) for e in scheme}
    return [by_name.get(name,"") for name in ["lt1","dk1","lt2","dk2","accent1","accent2","accent3","accent4","accent5","accent6","hlink","folHlink"]]

def color_hex(color, themes):
    if color.type=="rgb":raw=color.rgb[-6:]
    elif color.type=="theme":raw=themes[color.theme] if color.theme<len(themes) else ""
    elif color.type=="indexed":raw=COLOR_INDEX[color.indexed][-6:] if color.indexed<len(COLOR_INDEX) else ""
    else:return None
    if not re.fullmatch(r"[0-9a-fA-F]{6}",raw):return None
    rgb=[int(raw[i:i+2],16) for i in (0,2,4)]
    if color.tint:
        import colorsys
        h,l,s=colorsys.rgb_to_hls(*(x/255 for x in rgb))
        l=l*(1+color.tint) if color.tint<0 else l*(1-color.tint)+color.tint
        rgb=[round(x*255) for x in colorsys.hls_to_rgb(h,l,s)]
    return "#"+"".join(f"{x:02X}" for x in rgb)

def read_xlsx(content):
    check_archive(content)
    try:
        book=load_workbook(BytesIO(content),read_only=False,data_only=False,keep_links=False)
        if len(book.worksheets)>20:raise ValidationError("Se admiten hasta 20 hojas.")
        themes=theme_colors(book);sheets=[]
        for sheet in book.worksheets:
            warnings=[]
            if sheet.merged_cells.ranges:warnings.append("Hay celdas fusionadas: confirmá manualmente filas y fechas.")
            if sheet.conditional_formatting:warnings.append("Hay formato condicional: sus colores requieren revisión manual.")
            cells=[]
            # openpyxl 3.1.5 fijado: iter_rows crea celdas vacías en todo el rectángulo.
            for c in sorted(sheet._cells.values(),key=lambda c:(c.row,c.column)):
                note=c.comment.text if c.comment else ""
                color=color_hex(c.fill.fgColor,themes) if c.fill.patternType=="solid" else None
                if c.value is None and not note and not color:continue
                cells.append({"ref":c.coordinate,"row":c.row,"col":c.column,"value":label(c.value),
                    "note":note,"author":c.comment.author if c.comment else "",
                    "color":color,"kind":c.data_type})
            sheets.append({"name":sheet.title,"cells":cells,"warnings":warnings})
        book.close()
        return {"format":"xlsx","sheets":sheets}
    except ValidationError:raise
    except Exception as exc:raise ValidationError("No se pudo leer el Excel. Comprobá su formato.") from exc

def suggest_configuration(source):
    configs=[]
    for sheet in source.get("sheets",[]):
        cells=sheet["cells"];heads={}
        for c in cells:
            if c["kind"]!="f" and c["value"]=="1":heads.setdefault(c["row"],[]).append(c["col"])
        by_position={(c["row"],c["col"]):c["value"] for c in cells}
        headers=[r for r,cols in heads.items() if any(by_position.get((r,col+1))=="2" for col in cols)]
        label_col=max(1,min((c["col"] for c in cells if c["value"]),default=1))
        units=[c["row"] for c in cells if c["col"]==label_col and c["row"] not in headers and re.fullmatch(r"\d+(?:[- ]?[a-zA-Z]+)?",c["value"])]
        match=re.search(r"(?i)(ene|feb|mar|abr|may|jun|jul|ago|sep|oct|nov|dic)[^0-9]*(\d{2,4})",sheet["name"])
        year=int(match[2]) if match else 2026
        if year<100:year+=2000
        configs.append({"name":sheet["name"],"year":year,"month":MONTHS[match[1].lower()] if match else 0,
            "label_col":label_col,"header_rows":headers,"unit_rows":sorted(set(units))})
    return configs
