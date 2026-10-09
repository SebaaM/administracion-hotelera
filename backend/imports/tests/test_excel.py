from django.test import SimpleTestCase
from rest_framework.exceptions import ValidationError
from io import BytesIO
from zipfile import ZipFile
from openpyxl import load_workbook
from openpyxl.styles import Color, PatternFill
from .fixtures import calendar

class ExcelTests(SimpleTestCase):
    def test_comments_and_colors_survive(self):
        from imports.readers import read_xlsx, suggest_configuration
        result=read_xlsx(calendar())
        cells={c["ref"]:c for c in result["sheets"][0]["cells"]}
        self.assertEqual(cells["K2"]["note"],"Llegada tarde\nPedido especial")
        self.assertEqual(cells["K2"]["color"],"#C6EFCE")
        self.assertEqual(cells["A2"]["note"],"Nota de cama")
        self.assertEqual(suggest_configuration(result)[0]["unit_rows"],[2])
    def test_theme_and_numeric_labels(self):
        from imports.readers import read_xlsx
        book=load_workbook(BytesIO(calendar()))
        book.active["A3"]=10.0
        book.active["K2"].fill=PatternFill("solid",fgColor=Color(theme=4))
        out=BytesIO();book.save(out)
        cells={c["ref"]:c for c in read_xlsx(out.getvalue())["sheets"][0]["cells"]}
        self.assertEqual(cells["A3"]["value"],"10")
        self.assertEqual(cells["K2"]["color"],"#4F81BD")
    def test_invalid_and_macro_files_rejected(self):
        from imports.readers import read_xlsx
        for data in (b"not an Excel",b"x"*(5*1024*1024+1)):
            with self.assertRaises(ValidationError):read_xlsx(data)
        output=BytesIO()
        with ZipFile(output,"w") as z:z.writestr("xl/vbaProject.bin",b"x")
        with self.assertRaises(ValidationError):read_xlsx(output.getvalue())
    def test_declared_enormous_dimensions_rejected(self):
        from imports.readers import read_xlsx
        inp=ZipFile(BytesIO(calendar()));out=BytesIO()
        with ZipFile(out,"w") as z:
            for item in inp.infolist():
                data=inp.read(item.filename)
                if item.filename=="xl/worksheets/sheet1.xml":
                    data=data.replace(b'ref="A1:AF2"',b'ref="A1:XFD1048576"')
                z.writestr(item.filename,data)
        with self.assertRaises(ValidationError):read_xlsx(out.getvalue())
