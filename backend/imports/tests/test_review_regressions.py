from io import BytesIO
from zipfile import ZipFile
from datetime import timedelta
from unittest.mock import patch
from django.test import SimpleTestCase,TestCase
from django.utils import timezone
from rest_framework.exceptions import ValidationError
from openpyxl import Workbook
from imports.readers import check_archive,read_xlsx,LIMIT
from imports.csv_io import read_csv,write_csv
from imports.services import save_decisions
from imports.apply import apply_draft
from pms.models import Room,Unit,Reservation,Allocation,LedgerEntry
from .fixtures import calendar,source_candidate
from . import test_preview as setup

def edited_archive(transform):
    out=BytesIO()
    with ZipFile(BytesIO(calendar())) as src,ZipFile(out,"w") as dst:
        for item in src.infolist():
            dst.writestr(item.filename,transform(item.filename,src.read(item.filename)))
    return out.getvalue()

class FileReviewTests(SimpleTestCase):
    def test_expanding_ranges_rejected_before_openpyxl(self):
        for tag in ('<mergeCells count="1"><mergeCell ref="A1:XFD1048576"/></mergeCells>',
                    '<mergeCells count="1"><mergeCell ref="A1:OJ5000"/></mergeCells>',
                    '<hyperlinks><hyperlink ref="A1:XFD1048576" location="A1"/></hyperlinks>'):
            with self.subTest(tag=tag):
                content=edited_archive(lambda name,data:data.replace(b"</worksheet>",tag.encode()+b"</worksheet>") if name=="xl/worksheets/sheet1.xml" else data)
                with self.assertRaises(ValidationError):check_archive(content)
    def test_comment_coordinates_rejected_before_openpyxl(self):
        content=edited_archive(lambda name,data:data.replace(b'ref="K2"',b'ref="XFD1048576"') if "comment" in name and name.endswith(".xml") else data)
        with self.assertRaises(ValidationError):check_archive(content)
    def test_sparse_sheet_does_not_expand_its_bounding_rectangle(self):
        book=Workbook();book.active["A1"]="Etiqueta";book.active.cell(5000,400,"Dato ficticio")
        out=BytesIO();book.save(out)
        # La prueba impide materializar dos millones de celdas para medir dos datos.
        with patch("openpyxl.worksheet.worksheet.Worksheet.iter_rows",side_effect=AssertionError("Recorrido denso")):
            result=read_xlsx(out.getvalue())
        self.assertEqual(len(result["sheets"][0]["cells"]),2)
    def test_csv_metadata_is_rejected_before_a_draft_is_created(self):
        for extra in ({"references":["invalid"]},
                      {"notes":[{"text":"Nota ficticia","refs":["invalid"]}]},
                      {"colors":[{"hex":"#C6EFCE","meaning":{"invalid":True}}]},
                      {"colors":[{"hex":"invalid"}]}):
            with self.subTest(extra=extra):
                row=source_candidate(source_id="f03c4ea3-eaf9-4fab-9ac5-14feffecf776",**extra)
                with self.assertRaises(ValidationError):read_csv(write_csv([row]))
    def test_extended_notes_roundtrip_within_file_limit(self):
        notes=[{"text":str(index)+"x"*16000,"refs":[{"sheet":"AGO-26","cell":"K2"}]} for index in range(10)]
        row=source_candidate(source_id="f03c4ea3-eaf9-4fab-9ac5-14feffecf776",notes=notes)
        content=write_csv([row]);self.assertLess(len(content),LIMIT)
        self.assertEqual(read_csv(content)[0]["notes"],notes)

class OperationalReviewTests(TestCase):
    def setUp(self):setup.PreviewTests.setUp(self)
    def existing(self,guests=1,unit=None):
        start=timezone.localdate()+timedelta(days=2);end=start+timedelta(days=3)
        reservation=Reservation.objects.create(guest="Persona ficticia",start=start,end=end,guests=guests)
        Allocation.objects.create(reservation=reservation,unit=unit or self.unit,start=start,end=end,rate=100)
        return reservation
    def prepare(self,reservation,status="CONFIRMED",unit=None):
        draft=setup.PreviewTests.draft(self)
        draft.source["candidates"]=[source_candidate(start=str(reservation.start),end=str(reservation.end))]
        draft.save()
        decisions=setup.PreviewTests.reviewed(self)
        decisions["rows"]["AGO-26:2:10"].update(reservation_id=reservation.pk,status=status,unit_id=(unit or self.unit).pk)
        result=save_decisions(draft.pk,1,decisions,self.user);draft.refresh_from_db()
        return draft,result
    def test_import_cannot_cancel_reservation_with_payments(self):
        reservation=self.existing()
        LedgerEntry.objects.create(reservation=reservation,created_by=self.user,kind="CHARGE",amount=100,description="Cargo ficticio")
        LedgerEntry.objects.create(reservation=reservation,created_by=self.user,kind="PAYMENT",amount=50,description="Cobro ficticio")
        draft,result=self.prepare(reservation,"CANCELLED")
        self.assertFalse(result["can_apply"])
        with self.assertRaises(ValidationError):apply_draft(draft.pk,draft.revision,self.user)
        reservation.refresh_from_db();self.assertEqual(reservation.status,"CONFIRMED")
        self.assertEqual(LedgerEntry.objects.count(),2)
    def test_import_cannot_move_four_guests_to_one_bed(self):
        private=Room.objects.create(name="Privada",kind="PRIVATE",capacity=4)
        whole=Unit.objects.create(room=private,name="Completa")
        reservation=self.existing(4,whole)
        draft,result=self.prepare(reservation,unit=self.unit)
        self.assertFalse(result["can_apply"])
        with self.assertRaises(ValidationError):apply_draft(draft.pk,draft.revision,self.user)
        self.assertEqual(reservation.allocations.get().unit_id,whole.pk)
    def test_capacity_is_revalidated_when_applying(self):
        private=Room.objects.create(name="Privada",kind="PRIVATE",capacity=4)
        whole=Unit.objects.create(room=private,name="Completa")
        target_room=Room.objects.create(name="Destino",kind="PRIVATE",capacity=4)
        target=Unit.objects.create(room=target_room,name="Completa")
        reservation=self.existing(4,whole)
        draft,result=self.prepare(reservation,unit=target)
        self.assertTrue(result["can_apply"])
        target_room.capacity=1;target_room.save()
        with self.assertRaises(ValidationError):apply_draft(draft.pk,draft.revision,self.user)

class HeaderWorkReviewTests(SimpleTestCase):
    def test_many_numeric_cells_do_not_require_quadratic_scans(self):
        from imports.readers import suggest_configuration
        class BoundedCells(list):
            visits=0
            def __iter__(self):
                for cell in super().__iter__():
                    self.visits+=1
                    if self.visits>8*len(self):raise AssertionError("La detección excede el trabajo lineal permitido")
                    yield cell
        cells=BoundedCells({"row":row,"col":col,"kind":"n","value":"1"} for row in range(1,51) for col in range(1,21))
        result=suggest_configuration({"sheets":[{"name":"AGO-26","cells":cells}]})
        self.assertEqual(result[0]["header_rows"],[])
