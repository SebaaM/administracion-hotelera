from io import BytesIO
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
def calendar():
    book=Workbook()
    sheet=book.active
    sheet.title="AGO-26"
    sheet["A2"]="2-a"
    for day in range(1,32): sheet.cell(1,day+1,day)
    for day in (10,11,12):
        cell=sheet.cell(2,day+1,"Persona ficticia")
        cell.fill=PatternFill("solid",fgColor="C6EFCE")
    sheet["K2"].comment=Comment("Llegada tarde\nPedido especial","Prueba")
    sheet["N2"].fill=PatternFill("solid",fgColor="FF0000")
    sheet["A2"].comment=Comment("Nota de cama","Prueba")
    output=BytesIO(); book.save(output)
    return output.getvalue()
def source_candidate(**extra):
    return {"key":"AGO-26:2:10","source_id":"","source_unit":"2-a","guest":"Persona ficticia",
        "start":"2026-08-10","end":"2026-08-13","notes":[{"text":"Llegada tarde","refs":[{"sheet":"AGO-26","cell":"K2"}]}],
        "colors":[{"hex":"#C6EFCE","sheet":"AGO-26","cell":"K2","meaning":"Pagado"}],
        "references":[{"sheet":"AGO-26","cell":"K2"}],"warnings":[],**extra}
